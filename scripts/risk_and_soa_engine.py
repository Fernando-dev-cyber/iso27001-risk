import json
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Lista extendida de Controles ISO/IEC 27001:2022 (Anexo A)
ANNEX_A_CONTROLS = [
    # --- 5. Controles Organizacionales ---
    {"control_id": "A.5.1", "theme": "Organizational", "name": "Policies for information security", "applicable": "Aplica", "status": "Implemented", "justification": "Política de Seguridad publicada y revisada anualmente."},
    {"control_id": "A.5.7", "theme": "Organizational", "name": "Threat intelligence", "applicable": "Aplica", "status": "Implemented", "justification": "Integración de feeds de reputación de IPs/Hashes vía ioc-enricher (VT/AbuseIPDB)."},
    {"control_id": "A.5.15", "theme": "Organizational", "name": "Access control", "applicable": "Aplica", "status": "Implemented", "justification": "Política RBAC estricta y principio de menor privilegio en servicios cloud y locales."},
    {"control_id": "A.5.23", "theme": "Organizational", "name": "Information security for use of cloud services", "applicable": "Aplica", "status": "Implemented", "justification": "Monitoreo continuo y alertas en GCP IAM ante modificaciones no autorizadas."},
    {"control_id": "A.5.24", "theme": "Organizational", "name": "Information security incident management planning", "applicable": "Aplica", "status": "Implemented", "justification": "Plan de respuesta a incidentes documentado con playbooks de triage."},

    # --- 6. Controles de Personas ---
    {"control_id": "A.6.3", "theme": "People", "name": "Information security awareness, education and training", "applicable": "Aplica", "status": "Implemented", "justification": "Capacitación obligatoria trimestral en concientización sobre Phishing e Ingeniería Social."},
    {"control_id": "A.6.8", "theme": "People", "name": "Information security event reporting", "applicable": "Aplica", "status": "Implemented", "justification": "Canal de reporte de incidentes habilitado en Help Desk para usuarios finales."},

    # --- 7. Controles Físicos ---
    {"control_id": "A.7.1", "theme": "Physical", "name": "Physical security perimeters", "applicable": "No Aplica", "status": "Not Applicable", "justification": "Empresa 100% Cloud-native en GCP. La seguridad física es responsabilidad delegada del CSP."},
    {"control_id": "A.7.4", "theme": "Physical", "name": "Physical security monitoring", "applicable": "No Aplica", "status": "Not Applicable", "justification": "Sin centros de datos físicos propios. Infraestructura totalmente serverless/cloud."},

    # --- 8. Controles Tecnológicos ---
    {"control_id": "A.8.7", "theme": "Technological", "name": "Protection against malware", "applicable": "Aplica", "status": "Implemented", "justification": "Despliegue de agentes EDR/Antivirus y recolección volátil con FastTriage.ps1."},
    {"control_id": "A.8.8", "theme": "Technological", "name": "Management of technical vulnerabilities", "applicable": "Aplica", "status": "Implemented", "justification": "Auditoría periódica de superficie de ataque con Nmap Sentinel y Discord Alerting."},
    {"control_id": "A.8.12", "theme": "Technological", "name": "Data leakage prevention (DLP)", "applicable": "Aplica", "status": "In Progress", "justification": "Reglas DLP en fase de pilotaje para endpoints y repositorios de código."},
    {"control_id": "A.8.16", "theme": "Technological", "name": "Monitoring activities", "applicable": "Aplica", "status": "Implemented", "justification": "Ingesta de Cloud Audit Logs en tiempo real y triaje de tráfico en Wireshark."},
    {"control_id": "A.8.20", "theme": "Technological", "name": "Network security", "applicable": "Aplica", "status": "Implemented", "justification": "Segmentación de red VPC en GCP y aislamiento de tráfico sospechoso."},
    {"control_id": "A.8.28", "theme": "Technological", "name": "Secure coding", "applicable": "Aplica", "status": "In Progress", "justification": "Escaneo estático de código (SAST) y prevención de fuga de credenciales en Git."}
]

def calculate_risk_level(prob, imp):
    score = prob * imp
    if score >= 16:
        return score, "CRÍTICO"
    elif score >= 10:
        return score, "ALTO"
    elif score >= 5:
        return score, "MEDIO"
    else:
        return score, "BAJO"

def process_iso27001_data():
    json_path = os.path.join("data", "asset_inventory.json")
    if not os.path.exists(json_path):
        print("[!] Error: No se encontró data/asset_inventory.json")
        return

    with open(json_path, "r", encoding="utf-8") as f:
        assets = json.load(f)

    # Crear Libro de Trabajo nativo openpyxl
    wb = Workbook()
    
    # ---------------------------------------------------------
    # TAB 1: Risk Assessment
    # ---------------------------------------------------------
    ws_risk = wb.active
    ws_risk.title = "Risk_Assessment"
    
    risk_headers = [
        "ID Activo", "Nombre Activo", "Categoría", "Propietario",
        "Amenaza Principal", "Vulnerabilidad", "Probabilidad (1-5)",
        "Impacto (1-5)", "Riesgo Inherente", "Nivel Riesgo Inherente",
        "Estrategia Tratamiento", "Nivel Riesgo Residual"
    ]
    ws_risk.append(risk_headers)

    for ast in assets:
        prob = ast["probability"]
        imp = ast["impact"]
        score, level = calculate_risk_level(prob, imp)
        
        res_prob = max(1, prob - 2)
        res_imp = max(1, imp - 1)
        _, res_level = calculate_risk_level(res_prob, res_imp)

        ws_risk.append([
            ast["asset_id"], ast["asset_name"], ast["category"], ast["owner"],
            ast["threat"], ast["vulnerability"], prob, imp, score, level,
            "Mitigar / Reducir", res_level
        ])

    # ---------------------------------------------------------
    # TAB 2: Statement of Applicability (SoA)
    # ---------------------------------------------------------
    ws_soa = wb.create_sheet(title="Statement_of_Applicability")
    
    soa_headers = [
        "ID Control ISO", "Dominio / Tema", "Nombre del Control",
        "Aplica (SoA)", "Estado Implementación", "Justificación / Evidencia Técnica"
    ]
    ws_soa.append(soa_headers)

    for ctrl in ANNEX_A_CONTROLS:
        ws_soa.append([
            ctrl["control_id"], ctrl["theme"], ctrl["name"],
            ctrl["applicable"], ctrl["status"], ctrl["justification"]
        ])

    # ---------------------------------------------------------
    # ESTILOS PROFESIONALES DE AUDITORÍA
    # ---------------------------------------------------------
    HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Azul Marino
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    RISK_COLORS = {
        "CRÍTICO": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"), # Rojo
        "ALTO": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),    # Naranja
        "MEDIO": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),   # Amarillo
        "BAJO": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")     # Verde
    }
    
    STATUS_COLORS = {
        "Implemented": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "In Progress": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "Not Applicable": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    }

    THIN_BORDER = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for sheet in wb.worksheets:
        sheet.views.sheetView[0].showGridLines = True

        # Formatear encabezados
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Formatear filas de datos
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center")

                # Colores de Nivel de Riesgo
                if str(cell.value) in RISK_COLORS:
                    cell.fill = RISK_COLORS[str(cell.value)]
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Colores de Estado de Control
                if str(cell.value) in STATUS_COLORS:
                    cell.fill = STATUS_COLORS[str(cell.value)]
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # Ajuste dinámico de ancho de columnas
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    # ---------------------------------------------------------
    # GUARDAR ARCHIVO
    # ---------------------------------------------------------
    os.makedirs("output", exist_ok=True)
    output_excel = os.path.join("output", "ISO27001_Master_Register.xlsx")
    
    # Si el archivo está abierto en Excel, evitamos un PermissionError
    try:
        wb.save(output_excel)
        print(f"[+] Archivo Excel Maestro generado y formateado con éxito en: {output_excel}")
    except PermissionError:
        print("[!] Error: Cierra el archivo 'ISO27001_Master_Register.xlsx' en Microsoft Excel antes de volver a ejecutar el script.")

if __name__ == "__main__":
    process_iso27001_data()